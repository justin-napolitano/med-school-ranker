from __future__ import annotations

import csv
import json
import re
import zipfile
from html import unescape
from pathlib import Path

from openpyxl import load_workbook

from med_school_ranker import bundle, rankings, reviewer_state, site as site_builder, validation, workbook
from med_school_ranker.paths import ROOT


SCHOOL_MASTER_FIELDS = [
    "school_id",
    "school_name",
    "degree_type",
    "city",
    "state",
    "state_abbrev",
    "active_in_universe",
    "manual_exclusion_flag",
    "exclusion_reason",
    "admissions_mcat_fit_score",
    "attendance_prestige_score",
    "regret_index_score",
]


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def default_school_rows() -> list[dict[str, str]]:
    return [
        {
            "school_id": "school_one",
            "school_name": "School One",
            "degree_type": "MD",
            "city": "Boston",
            "state": "Massachusetts",
            "state_abbrev": "MA",
            "active_in_universe": "TRUE",
            "manual_exclusion_flag": "FALSE",
            "exclusion_reason": "",
            "admissions_mcat_fit_score": "8",
            "attendance_prestige_score": "9",
            "regret_index_score": "7",
        },
        {
            "school_id": "school_two",
            "school_name": "School Two",
            "degree_type": "DO",
            "city": "Chicago",
            "state": "Illinois",
            "state_abbrev": "IL",
            "active_in_universe": "TRUE",
            "manual_exclusion_flag": "FALSE",
            "exclusion_reason": "",
            "admissions_mcat_fit_score": "4",
            "attendance_prestige_score": "5",
            "regret_index_score": "3",
        },
    ]


def write_minimal_project(
    root: Path,
    school_rows: list[dict[str, str]] | None = None,
    partner_rows: list[dict[str, str]] | None = None,
) -> None:
    rows = school_rows or default_school_rows()
    write_csv(root / "data/school_master.csv", SCHOOL_MASTER_FIELDS, rows)
    write_csv(
        root / "data/user_preferences.csv",
        ["score_group", "column_name", "weight", "notes"],
        [
            {"score_group": "Admissions Score", "column_name": "admissions_mcat_fit_score", "weight": "1", "notes": ""},
            {"score_group": "Attendance Score", "column_name": "attendance_prestige_score", "weight": "1", "notes": ""},
            {"score_group": "Overall School Value", "column_name": "admissions_score", "weight": "1", "notes": ""},
            {"score_group": "Overall School Value", "column_name": "attendance_score", "weight": "1", "notes": ""},
            {"score_group": "Overall School Value", "column_name": "regret_index_score", "weight": "1", "notes": ""},
        ],
    )
    write_csv(
        root / "data/scenario_weights.csv",
        ["scenario", "column_name", "weight", "notes"],
        [{"scenario": "Balanced", "column_name": "overall_school_value", "weight": "1", "notes": ""}],
    )
    write_csv(
        root / "data/applicant_profiles.csv",
        validation.APPLICANT_PROFILE_COLUMNS,
        [
            {
                "applicant_profile_id": "template_profile",
                "profile_name": "Template Profile",
                "is_default_profile": "TRUE",
                "mcat_total": "",
                "overall_gpa": "",
                "science_gpa": "",
                "state_of_residence": "",
                "preferred_setting": "urban",
                "urban_preference_score": "8",
                "rural_tolerance_score": "5",
                "specialty_interest": "emergency medicine",
                "target_application_count": "25",
                "cost_weight_level": "low",
                "hidden_curriculum_weight_level": "low",
                "specialty_weight_level": "low",
                "notes": "Template only.",
            }
        ],
    )
    write_csv(root / "data/reference/aamc_mcat_gpa_acceptance_grid.csv", validation.AAMC_MCAT_GPA_GRID_COLUMNS, [])
    write_csv(root / "data/normalized/admissions_stats.csv", validation.ADMISSIONS_STATS_COLUMNS, [])
    write_csv(root / "data/normalized/cost_and_debt.csv", validation.COST_AND_DEBT_COLUMNS, [])
    write_csv(root / "data/normalized/admissions_policies.csv", validation.ADMISSIONS_POLICIES_COLUMNS, [])
    write_csv(root / "data/normalized/letter_requirements.csv", validation.LETTER_REQUIREMENTS_COLUMNS, [])
    write_csv(
        root / "data/manual/admissions_source_queue.csv",
        validation.ADMISSIONS_SOURCE_QUEUE_COLUMNS,
        [
            {
                "school_id": row["school_id"],
                "school_name": row["school_name"],
                "degree_type": row["degree_type"],
                "candidate_source_url": "",
                "source_status": "",
                "extraction_status": "",
                "review_status": "",
                "last_checked": "",
                "notes": "",
            }
            for row in rows
        ],
    )
    write_csv(root / "data/manual/source_review_queue.csv", validation.SOURCE_REVIEW_QUEUE_COLUMNS, [])
    write_csv(
        root / "data/manual/partner_inputs.csv",
        validation.PARTNER_INPUT_COLUMNS,
        partner_rows
        or [
            {
                "applicant_profile_id": "template_profile",
                "school_id": row["school_id"],
                "school_name": row["school_name"],
                "could_live_here_4_years_score": "",
                "location_fit_score": "",
                "culture_fit_score": "",
                "regret_index_score": "",
                "hard_no_flag": "",
                "hard_no_reason": "",
                "partner_notes": "",
            }
            for row in rows
        ],
    )
    write_csv(root / "data/manual/school_visibility.csv", validation.SCHOOL_VISIBILITY_COLUMNS, [])
    write_csv(root / "data/manual/school_dossiers.csv", validation.SCHOOL_DOSSIER_COLUMNS, [])
    write_csv(root / "data/final_application_list.csv", ["school_id", "school_name", "why_kept", "why_cut"], [])
    write_csv(
        root / "data/manual/source_match_overrides.csv",
        [
            "override_id",
            "source_table",
            "source_row_number",
            "source_school_name",
            "source_state",
            "source_degree_type",
            "override_action",
            "school_id",
            "school_name",
            "match_status",
            "reviewed_by",
            "reviewed_date",
            "review_notes",
        ],
        [],
    )
    write_csv(
        root / "data/project_subplans.csv",
        [
            "plan_name",
            "status",
            "priority",
            "phase",
            "current_decision_summary",
            "next_action",
            "doc_path",
        ],
        [
            {
                "plan_name": "Test Plan",
                "status": "active",
                "priority": "1",
                "phase": "test",
                "current_decision_summary": "Fixture row.",
                "next_action": "Run tests.",
                "doc_path": "docs/test.md",
            }
        ],
    )


def patch_ranking_paths(monkeypatch, root: Path) -> Path:
    out = root / "outputs"
    monkeypatch.setattr(rankings, "ROOT", root)
    monkeypatch.setattr(rankings, "OUT", out)
    monkeypatch.setattr(rankings, "MASTER_CSV", root / "data/school_master.csv")
    monkeypatch.setattr(rankings, "APPLICANT_PROFILES_CSV", root / "data/applicant_profiles.csv")
    monkeypatch.setattr(rankings, "PRIVATE_APPLICANT_PROFILES_CSV", root / "data/manual/private/applicant_profiles.local.csv")
    monkeypatch.setattr(rankings, "PREFERENCES_CSV", root / "data/user_preferences.csv")
    monkeypatch.setattr(rankings, "SCENARIO_WEIGHTS_CSV", root / "data/scenario_weights.csv")
    monkeypatch.setattr(rankings, "AAMC_MCAT_GPA_GRID_CSV", root / "data/reference/aamc_mcat_gpa_acceptance_grid.csv")
    monkeypatch.setattr(rankings, "ADMISSIONS_STATS_CSV", root / "data/normalized/admissions_stats.csv")
    monkeypatch.setattr(rankings, "COST_AND_DEBT_CSV", root / "data/normalized/cost_and_debt.csv")
    monkeypatch.setattr(rankings, "ADMISSIONS_POLICIES_CSV", root / "data/normalized/admissions_policies.csv")
    monkeypatch.setattr(rankings, "PARTNER_INPUTS_CSV", root / "data/manual/partner_inputs.csv")
    monkeypatch.setattr(rankings, "RANKINGS_CSV", out / "calculated_rankings.csv")
    monkeypatch.setattr(rankings, "SCORING_METHODOLOGY_CSV", out / "scoring_methodology.csv")
    monkeypatch.setattr(rankings, "SCORE_CONTRIBUTIONS_CSV", out / "score_contributions.csv")
    monkeypatch.setattr(rankings, "PRIVATE_OUT", out / "private")
    monkeypatch.setattr(rankings, "PRIVATE_RANKINGS_CSV", out / "private/calculated_rankings.private.csv")
    monkeypatch.setattr(rankings, "PRIVATE_SCORE_CONTRIBUTIONS_CSV", out / "private/score_contributions.private.csv")
    return out


def patch_site_paths(monkeypatch, root: Path) -> Path:
    out = root / "outputs"
    site_dir = out / "site"
    monkeypatch.setattr(site_builder, "ROOT", root)
    monkeypatch.setattr(site_builder, "OUT", out)
    monkeypatch.setattr(site_builder, "MASTER_CSV", root / "data/school_master.csv")
    monkeypatch.setattr(site_builder, "RANKINGS_CSV", out / "calculated_rankings.csv")
    monkeypatch.setattr(site_builder, "SCORING_METHODOLOGY_CSV", out / "scoring_methodology.csv")
    monkeypatch.setattr(site_builder, "SCORE_CONTRIBUTIONS_CSV", out / "score_contributions.csv")
    monkeypatch.setattr(site_builder, "APPLICANT_PROFILES_CSV", root / "data/applicant_profiles.csv")
    monkeypatch.setattr(site_builder, "PREFERENCES_CSV", root / "data/user_preferences.csv")
    monkeypatch.setattr(site_builder, "SCENARIO_WEIGHTS_CSV", root / "data/scenario_weights.csv")
    monkeypatch.setattr(site_builder, "AAMC_MCAT_GPA_GRID_CSV", root / "data/reference/aamc_mcat_gpa_acceptance_grid.csv")
    monkeypatch.setattr(site_builder, "ADMISSIONS_STATS_CSV", root / "data/normalized/admissions_stats.csv")
    monkeypatch.setattr(site_builder, "COST_AND_DEBT_CSV", root / "data/normalized/cost_and_debt.csv")
    monkeypatch.setattr(site_builder, "ADMISSIONS_POLICIES_CSV", root / "data/normalized/admissions_policies.csv")
    monkeypatch.setattr(site_builder, "LETTER_REQUIREMENTS_CSV", root / "data/normalized/letter_requirements.csv")
    monkeypatch.setattr(site_builder, "PARTNER_INPUTS_CSV", root / "data/manual/partner_inputs.csv")
    monkeypatch.setattr(site_builder, "SCHOOL_VISIBILITY_CSV", root / "data/manual/school_visibility.csv")
    monkeypatch.setattr(site_builder, "SCHOOL_DOSSIERS_CSV", root / "data/manual/school_dossiers.csv")
    monkeypatch.setattr(site_builder, "SOURCE_MATCH_OVERRIDES_CSV", root / "data/manual/source_match_overrides.csv")
    monkeypatch.setattr(site_builder, "SOURCE_REVIEW_QUEUE_CSV", root / "data/manual/source_review_queue.csv")
    monkeypatch.setattr(site_builder, "ADMISSIONS_SOURCE_QUEUE_CSV", root / "data/manual/admissions_source_queue.csv")
    monkeypatch.setattr(site_builder, "SOURCE_INTEGRATION_REPORT_CSV", out / "source_integration_report.csv")
    monkeypatch.setattr(site_builder, "SOURCE_MATCH_REVIEW_CSV", out / "source_match_review.csv")
    monkeypatch.setattr(site_builder, "ADMISSIONS_STATS_CANDIDATES_CSV", out / "admissions_stats_candidates.csv")
    monkeypatch.setattr(site_builder, "ADMISSIONS_STATS_CONFLICTS_CSV", out / "admissions_stats_conflicts.csv")
    monkeypatch.setattr(site_builder, "COST_AND_DEBT_CANDIDATES_CSV", out / "cost_and_debt_candidates.csv")
    monkeypatch.setattr(site_builder, "COST_AND_DEBT_REVIEW_CSV", out / "cost_and_debt_review.csv")
    monkeypatch.setattr(site_builder, "DATA_QUALITY_REPORT_CSV", out / "data_quality_report.csv")
    monkeypatch.setattr(site_builder, "SITE_DIR", site_dir)
    monkeypatch.setattr(site_builder, "SITE_INDEX_HTML", site_dir / "index.html")
    monkeypatch.setattr(site_builder, "SOURCE_TABLE_DIR", root / "data/source_tables")
    monkeypatch.setattr(site_builder, "SOURCE_DIFFS_DIR", out / "source_diffs")
    monkeypatch.setattr(site_builder, "RAW_AAMC_MSAR_DIR", root / "data/raw/aamc/msar_reports")
    monkeypatch.setattr(
        site_builder,
        "JSON_OUTPUTS",
        {
            "school_master": root / "data/school_master.csv",
            "calculated_rankings": out / "calculated_rankings.csv",
            "scoring_methodology": out / "scoring_methodology.csv",
            "score_contributions": out / "score_contributions.csv",
            "applicant_profiles": root / "data/applicant_profiles.csv",
            "aamc_mcat_gpa_grid": root / "data/reference/aamc_mcat_gpa_acceptance_grid.csv",
            "admissions_stats": root / "data/normalized/admissions_stats.csv",
            "cost_and_debt": root / "data/normalized/cost_and_debt.csv",
            "admissions_policies": root / "data/normalized/admissions_policies.csv",
            "letter_requirements": root / "data/normalized/letter_requirements.csv",
            "partner_inputs": root / "data/manual/partner_inputs.csv",
            "school_visibility": root / "data/manual/school_visibility.csv",
            "school_dossiers": root / "data/manual/school_dossiers.csv",
            "source_match_overrides": root / "data/manual/source_match_overrides.csv",
            "source_review_queue": root / "data/manual/source_review_queue.csv",
            "admissions_source_queue": root / "data/manual/admissions_source_queue.csv",
            "data_quality_report": out / "data_quality_report.csv",
            "source_integration_report": out / "source_integration_report.csv",
            "source_match_review": out / "source_match_review.csv",
            "admissions_stats_candidates": out / "admissions_stats_candidates.csv",
            "admissions_stats_conflicts": out / "admissions_stats_conflicts.csv",
            "cost_and_debt_candidates": out / "cost_and_debt_candidates.csv",
            "cost_and_debt_review": out / "cost_and_debt_review.csv",
            "project_subplans": root / "data/project_subplans.csv",
        },
    )
    return site_dir


def extract_embedded_payload(html_text: str) -> dict:
    embedded_match = re.search(
        r'<script type="application/json" id="site-data">(.*?)</script>',
        html_text,
        re.DOTALL,
    )
    assert embedded_match is not None
    return json.loads(unescape(embedded_match.group(1)))


def read_dict_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def test_ranking_generation_runs(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)

    output = rankings.build_rankings()

    with output.open(newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 2
    assert rows[0]["school_id"] == "school_one"
    assert rows[0]["overall_rank"] == "1"


def test_workbook_has_required_tabs(tmp_path, monkeypatch):
    required_tabs = {
        "Instructions",
        "Project Subplans",
        "School Master",
        "Applicant Profiles",
        "User Preferences",
        "Scenario Weights",
        "AAMC GPA MCAT Grid",
        "Admissions Stats",
        "Cost and Debt",
        "Admissions Policies",
        "Letter Requirements",
        "Admissions Source Queue",
        "Source Match Overrides",
        "Source Review Queue",
        "Partner Inputs",
        "School Visibility",
        "School Dossiers",
        "Calculated Rankings",
        "Final Application List",
        "Data Quality",
        "Source Integration Report",
        "Source Match Review",
        "Admissions Stats Candidates",
        "Admissions Stats Conflicts",
        "Cost and Debt Candidates",
        "Cost and Debt Review",
        "Sources",
        "Field Definitions",
    }
    assert required_tabs.issubset({title for title, _ in workbook.SHEETS})

    temp_sheets = []
    for title, path in workbook.SHEETS:
        if path is None:
            temp_sheets.append((title, None))
            continue
        temp_path = tmp_path / path.relative_to(ROOT)
        write_csv(temp_path, ["id"], [{"id": title}])
        temp_sheets.append((title, temp_path))

    monkeypatch.setattr(workbook, "ROOT", tmp_path)
    monkeypatch.setattr(workbook, "OUT", tmp_path / "outputs")
    monkeypatch.setattr(workbook, "WORKBOOK_XLSX", tmp_path / "outputs/med_school_ranker.xlsx")
    monkeypatch.setattr(workbook, "SHEETS", temp_sheets)

    output = workbook.build_workbook()
    wb = load_workbook(output, read_only=True)
    assert required_tabs.issubset(set(wb.sheetnames))


def test_validation_catches_duplicate_school_id(tmp_path):
    rows = default_school_rows()
    rows[1]["school_id"] = rows[0]["school_id"]
    write_minimal_project(tmp_path, school_rows=rows)

    issues = validation.validate_project(tmp_path)

    assert any(issue.severity == "error" and issue.field == "school_id" for issue in issues)


def test_validation_catches_invalid_score(tmp_path):
    rows = default_school_rows()
    rows[0]["admissions_mcat_fit_score"] = "11"
    write_minimal_project(tmp_path, school_rows=rows)

    issues = validation.validate_project(tmp_path)

    assert any(issue.severity == "error" and issue.field == "admissions_mcat_fit_score" for issue in issues)


def test_validation_catches_exclusion_without_reason(tmp_path):
    rows = default_school_rows()
    rows[0]["manual_exclusion_flag"] = "TRUE"
    rows[0]["exclusion_reason"] = ""
    write_minimal_project(tmp_path, school_rows=rows)

    issues = validation.validate_project(tmp_path)

    assert any(issue.severity == "error" and issue.field == "exclusion_reason" for issue in issues)


def test_site_generation_writes_local_payload_and_json(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    site_dir = patch_site_paths(monkeypatch, tmp_path)

    output = site_builder.build_site()

    assert output == site_dir / "index.html"
    html_text = output.read_text()
    assert 'id="site-data"' in html_text

    embedded_payload = extract_embedded_payload(html_text)

    master_rows = site_builder.read_csv(tmp_path / "data/school_master.csv")
    assert embedded_payload["meta"]["active_school_count"] == len(master_rows)
    assert embedded_payload["meta"]["site_mode"] == "local_full"
    assert embedded_payload["routes"]["default"] == "#/rankings"
    assert embedded_payload["routes"]["admin"]
    assert all(school["school_slug"] for school in embedded_payload["schools"])
    assert all(school["profile_route"].startswith("#/schools/") for school in embedded_payload["schools"])
    assert all(school["school"].get("state") != "Puerto Rico" for school in embedded_payload["schools"])

    required_json = [
        "school_master.json",
        "calculated_rankings.json",
        "applicant_profiles.json",
        "aamc_mcat_gpa_grid.json",
        "admissions_stats.json",
        "cost_and_debt.json",
        "admissions_policies.json",
        "letter_requirements.json",
        "partner_inputs.json",
        "school_visibility.json",
        "school_dossiers.json",
        "source_match_overrides.json",
        "source_review_queue.json",
        "admissions_source_queue.json",
        "data_quality_report.json",
        "source_integration_report.json",
        "source_match_review.json",
        "admissions_stats_candidates.json",
        "admissions_stats_conflicts.json",
        "cost_and_debt_candidates.json",
        "cost_and_debt_review.json",
        "project_subplans.json",
        "source_status.json",
        "curated_lists.json",
        "school_profiles.json",
        "public_sources.json",
        "site_payload.json",
    ]
    for filename in required_json:
        json.loads((site_dir / "data" / filename).read_text())


def test_publish_safe_site_excludes_admin_source_review_and_private_payloads(tmp_path, monkeypatch):
    secret = "super_secret_partner_review_value"
    partner_row = {
        "applicant_profile_id": "template_profile",
        "school_id": "school_one",
        "school_name": "School One",
        "could_live_here_4_years_score": "8",
        "location_fit_score": "7",
        "culture_fit_score": "6",
        "regret_index_score": "2",
        "hard_no_flag": "FALSE",
        "hard_no_reason": "",
        "partner_notes": secret,
    }
    write_minimal_project(tmp_path, partner_rows=[partner_row])
    write_csv(
        tmp_path / "data/manual/admissions_source_queue.csv",
        validation.ADMISSIONS_SOURCE_QUEUE_COLUMNS,
        [
            {
                "school_id": "school_one",
                "school_name": "School One",
                "degree_type": "MD",
                "candidate_source_url": "https://example.edu/admissions",
                "source_status": "found",
                "extraction_status": "",
                "review_status": "",
                "last_checked": "",
                "notes": secret,
            }
        ],
    )
    source_review_row = {field: "" for field in validation.SOURCE_REVIEW_QUEUE_COLUMNS}
    for field in source_review_row:
        if "note" in field:
            source_review_row[field] = secret
    write_csv(tmp_path / "data/manual/source_review_queue.csv", validation.SOURCE_REVIEW_QUEUE_COLUMNS, [source_review_row])
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    site_dir = patch_site_paths(monkeypatch, tmp_path)

    output = site_builder.build_site(site_mode="publish_safe")
    html_text = output.read_text()
    embedded_payload = extract_embedded_payload(html_text)

    assert embedded_payload["meta"]["site_mode"] == "publish_safe"
    assert embedded_payload["routes"]["admin"] == []
    assert "partner_inputs" not in embedded_payload
    assert "school_visibility" not in embedded_payload
    assert "school_dossiers" not in embedded_payload
    assert "source_review_queue" not in embedded_payload
    assert "admissions_source_queue" not in embedded_payload
    assert "data_quality_report" not in embedded_payload
    assert "source_status" not in embedded_payload
    assert all("partner_input" not in school for school in embedded_payload["schools"])
    assert all("visibility" not in school for school in embedded_payload["schools"])
    assert all("dossier" not in school for school in embedded_payload["schools"])
    assert all("admissions_source" not in school for school in embedded_payload["schools"])
    assert secret not in html_text

    generated_names = {path.relative_to(site_dir).as_posix() for path in site_dir.rglob("*") if path.is_file()}
    assert "data/partner_inputs.json" not in generated_names
    assert "data/school_visibility.json" not in generated_names
    assert "data/school_dossiers.json" not in generated_names
    assert "data/source_review_queue.json" not in generated_names
    assert "data/admissions_source_queue.json" not in generated_names
    assert "data/source_status.json" not in generated_names
    for path in site_dir.rglob("*"):
        if path.is_file():
            assert secret not in path.read_text(errors="ignore")


def test_curated_list_readiness_metadata_uses_ready_partial_provisional_semantics(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    cost_row = {field: "" for field in validation.COST_AND_DEBT_COLUMNS}
    cost_row.update(
        {
            "school_id": "school_one",
            "school_name": "School One",
            "degree_type": "MD",
            "in_state_tuition_fees_insurance": "40000",
            "out_state_tuition_fees_insurance": "65000",
            "estimated_coa_in_state": "72000",
            "estimated_coa_out_state": "97000",
        }
    )
    stats_row = {field: "" for field in validation.ADMISSIONS_STATS_COLUMNS}
    stats_row.update(
        {
            "school_id": "school_one",
            "school_name": "School One",
            "degree_type": "MD",
            "published_mcat_average": "511",
            "published_gpa_average": "3.75",
            "published_mcat_band": "510-513",
            "published_gpa_band": "3.60-3.79",
            "aamc_acceptance_rate_band": "40-54.9%",
            "data_quality_band": "high",
        }
    )
    write_csv(tmp_path / "data/normalized/cost_and_debt.csv", validation.COST_AND_DEBT_COLUMNS, [cost_row])
    write_csv(tmp_path / "data/normalized/admissions_stats.csv", validation.ADMISSIONS_STATS_COLUMNS, [stats_row])
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    patch_site_paths(monkeypatch, tmp_path)

    payload = site_builder.build_site_payload()
    lists = {row["list_id"]: row for row in payload["curated_lists"]}

    assert lists["best_low_cost"]["readiness_label"] == "partial"
    assert lists["best_admissions_realism"]["readiness_label"] == "partial"
    assert lists["best_public_schools"]["readiness_label"] == "provisional"
    assert lists["best_private_schools"]["readiness_label"] == "provisional"
    assert lists["best_cities"]["readiness_label"] == "provisional"
    assert lists["best_culture_fit"]["readiness_label"] == "provisional"
    assert {"ready", "partial", "provisional"} == set(payload["copy"]["curated_list_readiness"])
    assert "school-specific acceptance probability" in payload["copy"]["aamc_grid_caveat"]
    assert all(row["slug"] and row["route"] == f"#/lists/{row['slug']}" for row in payload["curated_lists"])


def test_default_rankings_route_and_admin_route_are_separated(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    patch_site_paths(monkeypatch, tmp_path)

    output = site_builder.build_site()
    html_text = output.read_text()
    embedded_payload = extract_embedded_payload(html_text)

    assert embedded_payload["routes"]["default"] == "#/rankings"
    assert embedded_payload["routes"]["public"][0]["path"] == "#/rankings"
    assert any(route["path"] == "#/admin" for route in embedded_payload["routes"]["admin"])
    assert 'data-route="#/rankings"' in html_text
    assert 'data-route="#/admin"' in html_text
    assert 'data-view="dashboard"' not in html_text


def test_site_contains_local_visibility_dossier_and_research_workflows(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    patch_site_paths(monkeypatch, tmp_path)

    output = site_builder.build_site()
    html_text = output.read_text()
    embedded_payload = extract_embedded_payload(html_text)
    public_routes = {route["path"] for route in embedded_payload["routes"]["public"]}

    assert "#/dossiers" in public_routes
    assert "#/research" in public_routes
    assert 'section id="dossiers"' in html_text
    assert 'section id="research"' in html_text
    assert "function renderDossiers()" in html_text
    assert "function renderResearch()" in html_text
    assert "school_visibility_v1" in html_text
    assert "school_dossier_edits_v1" in html_text
    assert "school_visibility_export.csv" in html_text
    assert "school_dossier_edits_export.csv" in html_text


def test_site_payload_loads_persisted_reviewer_state(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    write_csv(
        tmp_path / "data/manual/school_visibility.csv",
        validation.SCHOOL_VISIBILITY_COLUMNS,
        [
            {
                "school_id": "school_one",
                "school_name": "School One",
                "visibility_state": "hidden",
                "visibility_reason": "Not a current fit",
                "hidden_at": "2026-06-05T12:00:00Z",
                "updated_at": "2026-06-05T12:00:00Z",
                "source": "test",
                "notes": "",
            }
        ],
    )
    write_csv(
        tmp_path / "data/manual/school_dossiers.csv",
        validation.SCHOOL_DOSSIER_COLUMNS,
        [
            {
                "school_id": "school_one",
                "school_name": "School One",
                "research_status": "skimmed",
                "interest_level": "high",
                "four_year_happiness": "8",
                "location_fit": "7",
                "culture_fit": "",
                "regret_index": "6",
                "hard_no_flag": "",
                "hard_no_reason": "",
                "application_decision_status": "considering",
                "notes": "Review note",
                "updated_at": "2026-06-05T12:00:00Z",
                "source": "test",
            }
        ],
    )
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    patch_site_paths(monkeypatch, tmp_path)

    payload = site_builder.build_site_payload()
    school_one = next(school for school in payload["schools"] if school["school"]["school_id"] == "school_one")

    assert payload["school_visibility"][0]["visibility_state"] == "hidden"
    assert payload["school_dossiers"][0]["research_status"] == "skimmed"
    assert school_one["visibility"]["visibility_reason"] == "Not a current fit"
    assert school_one["dossier"]["notes"] == "Review note"


def test_reviewer_state_imports_browser_exports(tmp_path):
    write_minimal_project(tmp_path)
    write_csv(
        tmp_path / "data/manual/school_visibility.csv",
        validation.SCHOOL_VISIBILITY_COLUMNS,
        [
            {
                "school_id": "school_two",
                "school_name": "School Two",
                "visibility_state": "hidden",
                "visibility_reason": "Old reason",
                "hidden_at": "2026-06-04T12:00:00Z",
                "updated_at": "2026-06-04T12:00:00Z",
                "source": "test",
                "notes": "",
            }
        ],
    )
    visibility_export = tmp_path / "school_visibility_export.csv"
    write_csv(
        visibility_export,
        [
            "export_schema_version",
            "exported_at",
            "school_id",
            "school_name",
            "visibility_state",
            "visibility_reason",
            "hidden_at",
            "updated_at",
            "source",
            "notes",
        ],
        [
            {
                "export_schema_version": "school_visibility_v1",
                "exported_at": "2026-06-05T12:00:00Z",
                "school_id": "school_one",
                "school_name": "School One",
                "visibility_state": "hidden",
                "visibility_reason": "Not a current fit",
                "hidden_at": "2026-06-05T12:00:00Z",
                "updated_at": "2026-06-05T12:00:00Z",
                "source": "browser_session",
                "notes": "",
            },
            {
                "export_schema_version": "school_visibility_v1",
                "exported_at": "2026-06-05T12:00:00Z",
                "school_id": "school_two",
                "school_name": "School Two",
                "visibility_state": "visible",
                "visibility_reason": "",
                "hidden_at": "2026-06-04T12:00:00Z",
                "updated_at": "2026-06-05T12:00:00Z",
                "source": "browser_session",
                "notes": "",
            },
        ],
    )
    dossier_export = tmp_path / "school_dossier_edits_export.csv"
    write_csv(
        dossier_export,
        [
            "export_schema_version",
            "exported_at",
            "school_id",
            "school_name",
            "research_status",
            "interest_level",
            "four_year_happiness",
            "location_fit",
            "culture_fit",
            "regret_index",
            "hard_no_flag",
            "hard_no_reason",
            "application_decision_status",
            "notes",
        ],
        [
            {
                "export_schema_version": "school_dossier_edits_v1",
                "exported_at": "2026-06-05T12:00:00Z",
                "school_id": "school_one",
                "school_name": "School One",
                "research_status": "skimmed",
                "interest_level": "high",
                "four_year_happiness": "8",
                "location_fit": "7",
                "culture_fit": "",
                "regret_index": "6",
                "hard_no_flag": "",
                "hard_no_reason": "",
                "application_decision_status": "considering",
                "notes": "Review note",
            }
        ],
    )

    visibility_count = reviewer_state.import_visibility_export(
        visibility_export,
        target_path=tmp_path / "data/manual/school_visibility.csv",
        master_path=tmp_path / "data/school_master.csv",
    )
    dossier_count = reviewer_state.import_dossier_export(
        dossier_export,
        target_path=tmp_path / "data/manual/school_dossiers.csv",
        master_path=tmp_path / "data/school_master.csv",
    )

    visibility_rows = {row["school_id"]: row for row in read_dict_rows(tmp_path / "data/manual/school_visibility.csv")}
    dossier_rows = {row["school_id"]: row for row in read_dict_rows(tmp_path / "data/manual/school_dossiers.csv")}

    assert visibility_count == 2
    assert visibility_rows["school_one"]["visibility_state"] == "hidden"
    assert visibility_rows["school_two"]["visibility_state"] == "visible"
    assert dossier_count == 1
    assert dossier_rows["school_one"]["notes"] == "Review note"
    assert dossier_rows["school_one"]["source"] == "browser_dossier_export"


def test_school_profile_routes_resolve_for_every_active_school(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    patch_site_paths(monkeypatch, tmp_path)

    payload = site_builder.build_site_payload()
    profiles_by_slug = {profile["school_slug"]: profile for profile in payload["school_profiles"]}
    school_slugs = [school["school_slug"] for school in payload["schools"]]

    assert len(profiles_by_slug) == len(payload["schools"])
    assert len(set(school_slugs)) == len(school_slugs)
    for school in payload["schools"]:
        slug = school["school_slug"]
        assert slug in profiles_by_slug
        assert school["profile_route"] == f"#/schools/{slug}"
        assert profiles_by_slug[slug]["route"] == school["profile_route"]


def test_site_output_does_not_copy_private_data(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    private_file = tmp_path / "data/manual/private/secret.csv"
    private_file.parent.mkdir(parents=True)
    private_file.write_text("secret\nsuper_secret_value\n")
    validation.validate_project(tmp_path)
    patch_ranking_paths(monkeypatch, tmp_path)
    rankings.build_rankings()
    site_dir = patch_site_paths(monkeypatch, tmp_path)

    site_builder.build_site()

    generated_files = [path for path in site_dir.rglob("*") if path.is_file()]
    assert all("private" not in path.relative_to(site_dir).parts for path in generated_files)
    assert all("super_secret_value" not in path.read_text(errors="ignore") for path in generated_files)


def test_upload_bundle_includes_site_and_excludes_private_data(tmp_path, monkeypatch):
    write_minimal_project(tmp_path)
    out = tmp_path / "outputs"
    site_dir = out / "site"
    write_csv(out / "calculated_rankings.csv", ["school_id"], [{"school_id": "school_one"}])
    write_csv(out / "data_quality_report.csv", validation.REPORT_COLUMNS, [])
    site_dir.mkdir(parents=True)
    (site_dir / "index.html").write_text("<!doctype html>")
    (site_dir / "data").mkdir()
    (site_dir / "data/site_payload.json").write_text("{}")
    workbook_path = out / "med_school_ranker.xlsx"
    workbook_path.write_text("workbook placeholder")
    private_file = tmp_path / "data/private/secret.csv"
    private_file.parent.mkdir(parents=True)
    private_file.write_text("secret\nsuper_secret_value\n")
    private_output = out / "private/calculated_rankings.private.csv"
    private_output.parent.mkdir(parents=True)
    private_output.write_text("secret\nsuper_secret_private_score\n")
    private_contribution_output = out / "private/score_contributions.private.csv"
    private_contribution_output.write_text("secret\nsuper_secret_private_contribution\n")

    monkeypatch.setattr(bundle, "ROOT", tmp_path)
    monkeypatch.setattr(bundle, "DATA", tmp_path / "data")
    monkeypatch.setattr(bundle, "OUT", out)
    monkeypatch.setattr(bundle, "SITE_DIR", site_dir)
    monkeypatch.setattr(bundle, "WORKBOOK_XLSX", workbook_path)
    monkeypatch.setattr(bundle, "UPLOAD_ZIP", tmp_path / "med-school-ranker-upload.zip")

    zip_path = bundle.build_upload_zip()

    with zipfile.ZipFile(zip_path) as archive:
        names = set(archive.namelist())
    assert "outputs/site/index.html" in names
    assert "outputs/site/data/site_payload.json" in names
    assert not any(name.startswith("data/private/") for name in names)
    assert not any(name.startswith("data/manual/private/") for name in names)
    assert not any(name.startswith("outputs/private/") for name in names)
