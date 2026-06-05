from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from med_school_ranker.paths import DATA, OUT, RANKINGS_CSV, ROOT, WORKBOOK_XLSX
from med_school_ranker.paths import (
    ADMISSIONS_POLICIES_CSV,
    ADMISSIONS_SOURCE_QUEUE_CSV,
    ADMISSIONS_STATS_CANDIDATES_CSV,
    ADMISSIONS_STATS_CONFLICTS_CSV,
    ADMISSIONS_STATS_CSV,
    APPLICANT_PROFILES_CSV,
    COST_AND_DEBT_CANDIDATES_CSV,
    COST_AND_DEBT_CSV,
    COST_AND_DEBT_REVIEW_CSV,
    DATA_QUALITY_REPORT_CSV,
    LETTER_REQUIREMENTS_CSV,
    PARTNER_INPUTS_CSV,
    SOURCE_INTEGRATION_REPORT_CSV,
    SOURCE_MATCH_OVERRIDES_CSV,
    SOURCE_MATCH_REVIEW_CSV,
    SOURCE_REVIEW_QUEUE_CSV,
)


SHEETS = [
    ("Instructions", None),
    ("Project Subplans", DATA / "project_subplans.csv"),
    ("School Master", DATA / "school_master.csv"),
    ("Applicant Profiles", APPLICANT_PROFILES_CSV),
    ("User Preferences", DATA / "user_preferences.csv"),
    ("Scenario Weights", DATA / "scenario_weights.csv"),
    ("Admissions Stats", ADMISSIONS_STATS_CSV),
    ("Cost and Debt", COST_AND_DEBT_CSV),
    ("Admissions Policies", ADMISSIONS_POLICIES_CSV),
    ("Letter Requirements", LETTER_REQUIREMENTS_CSV),
    ("Admissions Source Queue", ADMISSIONS_SOURCE_QUEUE_CSV),
    ("Source Match Overrides", SOURCE_MATCH_OVERRIDES_CSV),
    ("Source Review Queue", SOURCE_REVIEW_QUEUE_CSV),
    ("Partner Inputs", PARTNER_INPUTS_CSV),
    ("Calculated Rankings", RANKINGS_CSV),
    ("Final Application List", DATA / "final_application_list.csv"),
    ("Data Quality", DATA_QUALITY_REPORT_CSV),
    ("Source Integration Report", SOURCE_INTEGRATION_REPORT_CSV),
    ("Source Match Review", SOURCE_MATCH_REVIEW_CSV),
    ("Admissions Stats Candidates", ADMISSIONS_STATS_CANDIDATES_CSV),
    ("Admissions Stats Conflicts", ADMISSIONS_STATS_CONFLICTS_CSV),
    ("Cost and Debt Candidates", COST_AND_DEBT_CANDIDATES_CSV),
    ("Cost and Debt Review", COST_AND_DEBT_REVIEW_CSV),
    ("Sources", DATA / "sources.csv"),
    ("Field Definitions", DATA / "field_definitions.csv"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTLE_FILL = PatternFill("solid", fgColor="EAF2F8")
LINK_FONT = Font(color="0563C1", underline="single")


def read_rows(path: Path) -> List[List[str]]:
    with path.open(newline="") as f:
        return list(csv.reader(f))


def clean_sheet_name(name: str) -> str:
    return name[:31]


def as_number(value: str):
    if value is None:
        return ""
    text = value.strip()
    if not text:
        return ""
    if text in {"TRUE", "FALSE"}:
        return text == "TRUE"
    try:
        if text.startswith("0") and len(text) > 1 and not text.startswith("0."):
            return text
        number = float(text.replace(",", ""))
    except ValueError:
        return text
    if number.is_integer():
        return int(number)
    return number


def autofit_columns(ws, rows: Iterable[Iterable[str]], max_width: int = 48) -> None:
    widths = {}
    for row in rows:
        for idx, value in enumerate(row, start=1):
            text = str(value or "")
            widths[idx] = max(widths.get(idx, 0), min(len(text) + 2, max_width))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 10)


def style_header(row: Iterable[Cell]) -> None:
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_csv_sheet(wb: Workbook, title: str, path: Path) -> None:
    rows = read_rows(path)
    ws = wb.create_sheet(clean_sheet_name(title))
    for row in rows:
        ws.append([as_number(value) for value in row])

    if not rows:
        return

    style_header(ws[1])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = True
    ws.row_dimensions[1].height = 34

    table_ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
    table = Table(displayName=title.replace(" ", ""), ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    autofit_columns(ws, rows)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value
                cell.font = LINK_FONT


def add_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    rows = [
        ["Medical School Ranker", "Workbook generated from the CSV seed project."],
        ["Use", "Edit School Master scores and preferences, then regenerate rankings with uv."],
        ["Regenerate", "uv run med-school-build-all"],
        ["Primary tabs", "Project Subplans, School Master, Applicant Profiles, Admissions Stats, Cost and Debt, Admissions Policies, Letter Requirements, Source Review Queue, Partner Inputs, Calculated Rankings, Final Application List, Data Quality."],
        ["Important", "Calculated Rankings is generated output. Do not manually edit it as the source of truth."],
        ["Transparency", "Keep manual_exclusion_flag, exclusion_reason, source URLs, data_confidence, and last_verified_date current."],
        ["Score scale", "Use 1-10. Higher should always mean better fit or lower concern for the applicant."],
    ]
    for row in rows:
        ws.append(row)

    ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
    ws["B1"].font = Font(bold=True, size=12)
    for cell in ws["A"]:
        cell.fill = SUBTLE_FILL
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"


def build_workbook() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    for title, path in SHEETS:
        if path is None:
            add_instructions_sheet(wb)
        else:
            add_csv_sheet(wb, title, path)

    wb.save(WORKBOOK_XLSX)
    return WORKBOOK_XLSX


def main() -> None:
    output = build_workbook()
    print(f"Wrote {output.relative_to(ROOT)}")
